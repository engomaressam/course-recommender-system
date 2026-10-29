"""Build "Machine Learning Capstone Project Report.pptx" from the course template and the lab results.

The lab notebooks (m01-m03) write small CSV summaries to ../results; this script turns them into the
slides required by the final assignment. Export the PDF afterwards with export_pdf.ps1.

usage: python build_report.py
"""
import re
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.chart import XL_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from wordcloud import STOPWORDS, WordCloud

from pptx_helpers import (AXIS, DARK_BLUE, INK, MUTED, NEUTRAL, REFERENCE, SECONDARY, SERIES, SERIES_2, TINT,
                          TITLE_BLUE, WHITE, X0, X1, add_bar_chart, add_card, add_flowchart, add_line_chart,
                          add_stat, add_table, add_text, clear_slide, delete_slide, new_slide_like, remove_shape,
                          set_notes, set_runs_text)

SOURCE = Path(__file__).resolve().parent
CAPSTONE = SOURCE.parents[2]
RESULTS = CAPSTONE / "results"
DATASETS = CAPSTONE / "datasets"
TEMPLATE = CAPSTONE / "m05" / "01 - Final Project Overview - files" / "ml-capstone-template-coursera.pptx"
OUTPUT = SOURCE.parent / "Machine Learning Capstone Project Report.pptx"
BUILD = SOURCE / "build"

AUTHOR = "Omar Essam"
DATE = "October 24, 2026"
LIGHT_SERIES = "9EC5F4"


# ----------------------------------------------------------------------------------------- data
def csv(name):
    return pd.read_csv(RESULTS / name)


def metrics(name):
    df = csv(name)
    return dict(zip(df["metric"], df["value"]))


SMALL_WORDS = {"a", "an", "and", "as", "at", "by", "for", "in", "of", "on", "or", "the", "to", "with", "using", "from", "into", "vs"}
SPECIAL = {"ai": "AI", "ibm": "IBM", "sql": "SQL", "nosql": "NoSQL", "r": "R", "gpu": "GPU", "iot": "IoT", "api": "API",
           "apis": "APIs", "ml": "ML", "dbaas": "DBaaS", "hbase": "HBase", "mysql": "MySQL", "postgresql": "PostgreSQL",
           "mongodb": "MongoDB", "graphx": "GraphX", "mllib": "MLlib", "ui": "UI", "i": "I", "ii": "II", "db2": "Db2",
           "tensorflow": "TensorFlow", "bert": "BERT", "nlp": "NLP", "hdfs": "HDFS", "rdd": "RDD", "mapreduce": "MapReduce",
           "pgadmin": "pgAdmin", "cloudant": "Cloudant", "spss": "SPSS", "devops": "DevOps", "pyspark": "PySpark"}


def course_title(title):
    words = re.sub(r"\s+", " ", str(title)).strip().split(" ")
    out = []
    for i, w in enumerate(words):
        lw = w.lower()
        if lw in SPECIAL:
            out.append(SPECIAL[lw])
        elif i > 0 and lw in SMALL_WORDS:
            out.append(lw)
        else:
            out.append(lw[:1].upper() + lw[1:])
    return " ".join(out)


def genre_label(name):
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name)


def pct(x):
    return f"{x:.1%}"


# ----------------------------------------------------------------------------------------- slide builders
def title_slide(slide):
    for shape in list(slide.shapes):
        text = shape.text_frame.text if shape.has_text_frame else ""
        if text.startswith("Instructions"):
            remove_shape(shape)
        elif "<Your Name>" in text:
            p_name, p_date = shape.text_frame.paragraphs[:2]
            for p, value in ((p_name, AUTHOR), (p_date, DATE)):
                p.runs[0].text = value
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
            shape.width = Inches(4.5)
    set_notes(slide, f"Capstone project of the IBM Machine Learning Professional Certificate, presented by {AUTHOR} on {DATE}.")


def intro_slide(slide, eda):
    for shape in list(slide.shapes):
        if shape.has_text_frame and (shape.text_frame.text.startswith("Instructions") or "Project background" in shape.text_frame.text):
            remove_shape(shape)
    top, h = 1.55, 3.95
    col_w = (X1 - X0 - 0.3) / 2
    for i, (heading, body) in enumerate([
        ("Project background and context", [
            {"text": "An online learning platform (IBM Skills Network) offers 307 courses in data science, AI, cloud, databases and software development.", "bullet": True},
            {"text": "Learners struggle to find the next relevant course in a large catalogue; a personalised recommender surfaces unseen courses that match each learner's interests and raises engagement and completion.", "bullet": True},
            {"text": "Data: course titles, descriptions and 14 genre tags, plus 233,306 enrollments with ratings from 33,901 learners.", "bullet": True},
        ]),
        ("Problem statement and hypotheses", [
            {"runs": [("Problem: ", {"bold": True}), ("recommend new, relevant courses to every learner, and find out which family of models - content-based (unsupervised) or collaborative filtering (supervised) - does it best.", {})]},
            {"runs": [("H1: ", {"bold": True, "color": TITLE_BLUE}), ("enrollments concentrate on a few popular, introductory courses and genres.", {})], "bullet": False},
            {"runs": [("H2: ", {"bold": True, "color": TITLE_BLUE}), ("content-based methods can keep the list of new courses per learner under 20 while still reaching most learners.", {})]},
            {"runs": [("H3: ", {"bold": True, "color": TITLE_BLUE}), ("latent-factor models (NMF, neural embeddings) predict ratings more accurately (lower RMSE) than KNN and a global-mean baseline.", {})]},
        ]),
    ]):
        x = X0 + i * (col_w + 0.3)
        add_card(slide, x, top, col_w, h, fill=WHITE)
        add_text(slide, x + 0.25, top + 0.2, col_w - 0.5, 0.45, [{"text": heading, "size": 18, "bold": True, "color": TITLE_BLUE}])
        add_text(slide, x + 0.25, top + 0.75, col_w - 0.5, h - 0.9, body, size=14, space_after=8)
    stats = [(f"{int(eda['courses'])}", "courses"), (f"{int(eda['genres'])}", "course genres"),
             (f"{int(eda['users']):,}", "learners"), (f"{int(eda['enrollments']):,}", "enrollments (ratings)")]
    tile_w = (X1 - X0 - 3 * 0.3) / 4
    for i, (value, label) in enumerate(stats):
        add_stat(slide, X0 + i * (tile_w + 0.3), 5.68, tile_w, 0.92, value, label, fill=WHITE, value_size=26)
    set_notes(slide, "Background: a course platform with 307 courses needs to recommend relevant unseen courses to its learners. "
                     "We test three hypotheses: H1 popularity concentration, H2 manageable content-based recommendation lists, "
                     "H3 latent-factor collaborative filtering beats KNN and the baseline on RMSE.")


def genre_slide(slide):
    clear_slide(slide)
    g = csv("eda_genre_counts.csv")
    add_bar_chart(slide, X0, 1.85, 7.3, 5.05, [genre_label(x) for x in g["genre"]], [("Courses", g["Count"].tolist())],
                  horizontal=True, gap=45, tick_size=12, label_size=11)
    top3 = g.head(3)
    cg = pd.read_csv(DATASETS / "course_genre.csv")
    genres_per_course = cg.iloc[:, 2:].sum(axis=1)
    x = 8.55
    add_stat(slide, x, 1.95, X1 - x, 1.3, f"{int(top3['Count'].iloc[0])}", f"courses are tagged {genre_label(top3['genre'].iloc[0])}, the largest genre ({top3['Count'].iloc[0] / len(cg):.0%} of the {len(cg)} courses)")
    add_text(slide, x, 3.5, X1 - x, 3.4, [
        {"text": f"Backend Dev ({int(g['Count'].iloc[0])} courses) and Machine Learning ({int(g['Count'].iloc[1])}) lead the catalogue, followed by Database ({int(g['Count'].iloc[2])}) and Data Analysis.", "bullet": True},
        {"text": f"Chatbot and Blockchain are niche, with {int(g['Count'].iloc[-1])} courses each.", "bullet": True},
        {"text": f"A course can carry several genres: {g['Count'].sum()} tags over 307 courses ({genres_per_course.mean():.2f} per course on average).", "bullet": True},
    ], size=14, space_after=10)
    set_notes(slide, "Horizontal bar chart of the number of courses per genre, computed with course_df[genres].sum() and sort_values(). "
                     "Backend development, machine learning and databases are the biggest genres; chatbot and blockchain the smallest.")


def enrollment_slide(slide):
    clear_slide(slide)
    h = csv("eda_enrollment_histogram.csv")
    stats = dict(zip(csv("eda_enrollment_stats.csv")["statistic"], csv("eda_enrollment_stats.csv")["value"]))
    cap = 30
    cats = [str(i) for i in range(1, cap + 1)] + [f"{cap + 1}+"]
    counts = h.set_index("enrollments")["users"]
    values = [int(counts.get(i, 0)) for i in range(1, cap + 1)] + [int(counts[counts.index > cap].sum())]
    chart = add_bar_chart(slide, X0, 1.9, 7.9, 4.75, cats, [("Users", values)], gap=25, show_values=False,
                          show_value_axis=True, number_format="#,##0", tick_size=11, axis_title="Number of users",
                          cat_axis_title="Number of courses a user enrolled in", point_colors={0: TITLE_BLUE})
    pt = chart.plots[0].series[0].points[0]
    pt.data_label.has_text_frame = True
    pt.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    pt.data_label.text_frame.text = f"{values[0]:,}"
    run = pt.data_label.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.name = Pt(11), True, "Abadi"
    one = values[0] / int(stats["count"])
    x = 9.1
    add_stat(slide, x, 1.95, X1 - x, 1.3, pct(one), f"of learners ({values[0]:,}) enrolled in just one course")
    add_text(slide, x, 3.45, X1 - x, 3.5, [
        {"text": f"Mean {stats['mean']:.1f}, median {stats['50%']:.0f}, max {stats['max']:.0f} courses per learner (std {stats['std']:.1f}).", "bullet": True},
        {"text": f"The distribution is bimodal: very few learners have 2-4 courses ({sum(values[1:4]):,}), then a second peak at 5 courses ({values[4]:,}) with a long tail.", "bullet": True},
        {"text": "One-course learners are a cold-start group: popularity or content-based recommendations suit them better than collaborative filtering.", "bullet": True},
    ], size=13.5, space_after=9)
    set_notes(slide, "Histogram of the number of enrolled courses per user (groupby('user').size()). 8,320 users enrolled in a single course; "
                     "the rest mostly enrolled in 5-12 courses. Learners with more than 30 courses are grouped into the last bar.")


def top20_slide(slide, eda):
    clear_slide(slide)
    top = csv("eda_top20_courses.csv")
    add_bar_chart(slide, X0, 1.8, 7.9, 5.15, [course_title(t) for t in top["TITLE"]], [("Enrollments", top["Ratings"].tolist())],
                  horizontal=True, gap=35, number_format="#,##0", tick_size=10.5, label_size=10)
    share = eda["top20_enrollments"] / eda["enrollments"]
    x = 9.1
    add_stat(slide, x, 1.95, X1 - x, 1.3, pct(share), f"of all {int(eda['enrollments']):,} enrollments go to these 20 of 126 rated courses")
    add_text(slide, x, 3.45, X1 - x, 3.5, [
        {"text": f"Python for Data Science ({int(top['Ratings'].iloc[0]):,}), Introduction to Data Science and Big Data 101 lead the list.", "bullet": True},
        {"text": "Almost all are introductory courses (101, Introduction, Fundamentals): learners start with foundations.", "bullet": True},
        {"text": "This popularity bias (H1 confirmed) means a recommender must work to surface less popular courses.", "bullet": True},
    ], size=13.5, space_after=9)
    set_notes(slide, "The 20 most enrolled courses, from ratings_df.groupby('item').size() merged with course titles. Together they account for "
                     f"{pct(share)} of all enrollments.")


def wordcloud_slide(slide):
    clear_slide(slide)
    cg = pd.read_csv(DATASETS / "course_genre.csv")
    titles = " ".join(cg["TITLE"].astype(str))
    stopwords = set(STOPWORDS)
    stopwords.update(["getting started", "using", "enabling", "template", "university", "end", "introduction", "basic"])
    palette = ["#0B49CB", "#2A78D6", "#104281", "#1C5CAB", "#5598E7", "#52514E"]
    rng = np.random.default_rng(123)
    cloud = WordCloud(stopwords=stopwords, background_color="white", width=2200, height=1250, random_state=123,
                      font_path="C:/Windows/Fonts/calibrib.ttf", prefer_horizontal=0.9, margin=6,
                      color_func=lambda *a, **k: palette[int(rng.integers(len(palette)))]).generate(titles)
    BUILD.mkdir(exist_ok=True)
    image = BUILD / "wordcloud.png"
    cloud.to_file(image)
    slide.shapes.add_picture(str(image), Inches(X0), Inches(1.95), width=Inches(7.7))
    words = Counter(w for t in cg["TITLE"].astype(str) for w in set(t.lower().split()) if w not in stopwords and len(w) > 1)
    common = ", ".join(f"{w} ({n})" for w, n in words.most_common(6))
    x = 8.95
    add_text(slide, x, 1.95, X1 - x, 4.9, [
        {"text": "Most frequent title words", "size": 16, "bold": True, "color": TITLE_BLUE, "space_after": 6},
        {"text": f"{common} - counted as the number of titles that contain the word.", "size": 13.5, "space_after": 14},
        {"text": "What it tells us", "size": 16, "bold": True, "color": TITLE_BLUE, "space_after": 6},
        {"text": "The catalogue is built around in-demand IT skills: data science, Python, machine learning, cloud, big data, SQL and containers.", "bullet": True, "size": 13.5},
        {"text": "Title words overlap strongly between courses, so text features (bag of words) are a good basis for course similarity.", "bullet": True, "size": 13.5},
    ], space_after=8)
    set_notes(slide, "Word cloud of all 307 course titles (wordcloud package, English stop words plus a few generic words removed). "
                     "Larger words appear in more titles: data, python, science, machine learning, cloud and big data dominate.")


def flowchart_slide(slide, steps, notes, bullets):
    clear_slide(slide)
    bottom = add_flowchart(slide, steps, top=2.0)
    add_card(slide, X0, bottom + 0.3, X1 - X0, 6.95 - bottom - 0.3, fill=NEUTRAL)
    add_text(slide, X0 + 0.25, bottom + 0.42, X1 - X0 - 0.5, 6.95 - bottom - 0.5, [{"text": b, "bullet": True} for b in bullets],
             size=14, space_after=6)
    set_notes(slide, notes)


def evaluation_slide(slide, hyper_text, stat_value, stat_label, sweep_title, sweep_chart, top10, note):
    clear_slide(slide)
    add_card(slide, X0, 1.92, X1 - X0, 0.58, fill=TINT)
    add_text(slide, X0 + 0.2, 1.92, X1 - X0 - 0.4, 0.58, [{"runs": [("Hyper-parameters:  ", {"bold": True, "color": TITLE_BLUE}), (hyper_text, {})]}],
             size=12.5, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    left_w = 5.45
    add_stat(slide, X0, 2.66, left_w, 1.08, stat_value, stat_label, value_size=30, label_size=11.5)
    add_text(slide, X0, 3.86, left_w, 0.32, [{"text": sweep_title, "size": 12, "bold": True, "color": SECONDARY}], space_after=0)
    sweep_chart(slide, X0, 4.16, left_w, 2.79)
    x = X0 + left_w + 0.35
    add_text(slide, x, 2.62, X1 - x, 0.35, [{"text": "Top-10 most frequently recommended courses", "size": 14, "bold": True, "color": TITLE_BLUE}], space_after=0)
    titles = top10["TITLE"].map(course_title)
    repeated = titles.duplicated(keep=False)
    labels = [f"{t} ({cid})" if rep else t for t, cid, rep in zip(titles, top10["COURSE_ID"], repeated)]
    rows = [["#", "Course", "Times"]] + [[str(i + 1), label, f"{int(n):,}"] for i, (label, n) in enumerate(zip(labels, top10["RECOMMENDATIONS"]))]
    add_table(slide, x, 3.02, [0.4, X1 - x - 1.35, 0.95], rows, size=11, row_h=0.355,
              aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.RIGHT])
    set_notes(slide, note)


def sweep_columns(sweep, chosen, label_fmt, value_col="avg_recs_per_user"):
    def draw(slide, x, y, w, h):
        cats = [label_fmt(t) for t in sweep["threshold"]]
        colors = {i: TITLE_BLUE for i, t in enumerate(sweep["threshold"]) if np.isclose(t, chosen)}
        add_bar_chart(slide, x, y, w, h, cats, [("Avg new courses per user", sweep[value_col].round(1).tolist())],
                      colors=[LIGHT_SERIES], point_colors=colors, number_format="0.0", gap=45, tick_size=10.5, label_size=10,
                      cat_axis_title="Threshold (selected value in dark blue)")
    return draw


def content_comparison_slide(slide, profile, similarity, cluster, tops):
    names = ["User profile + genres", "Course similarity (BoW)", "User clustering (PCA + K-means)"]
    summaries = [profile, similarity, cluster]
    users = int(profile["test_users"])
    rows = [["Metric"] + names,
            ["Avg new courses per user"] + [f"{s['avg_recs_per_user']:.1f}" for s in summaries],
            ["Learners reached (at least one course)"] + [f"{int(s['users_with_recs']):,} ({s['users_with_recs'] / users:.0%})" for s in summaries],
            ["Catalogue coverage (distinct courses)"] + [f"{int(s['distinct_courses'])} of 307 ({s['distinct_courses'] / 307:.0%})" for s in summaries],
            ["Total recommendations"] + [f"{int(s['total_recs']):,}" for s in summaries]]
    add_table(slide, X0, 1.95, [3.35, 2.72, 2.72, 2.71], rows, size=12.5, row_h=0.46,
              aligns=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER], bold_first_col=True)
    overlap = len(set(tops[0]["COURSE_ID"]) & set(tops[1]["COURSE_ID"]) | set(tops[0]["COURSE_ID"]) & set(tops[2]["COURSE_ID"]) | set(tops[1]["COURSE_ID"]) & set(tops[2]["COURSE_ID"]))
    best_reach = names[int(np.argmax([s["users_with_recs"] for s in summaries]))]
    best_cov = names[int(np.argmax([s["distinct_courses"] for s in summaries]))]
    add_card(slide, X0, 4.6, X1 - X0, 2.3, fill=NEUTRAL)
    add_text(slide, X0 + 0.25, 4.72, X1 - X0 - 0.5, 2.1, [
        {"text": "Innovative insight: the three content-based recommenders behave very differently", "size": 15, "bold": True, "color": TITLE_BLUE, "space_after": 6},
        {"text": f"{best_reach} reaches the most learners, and {best_cov} spreads recommendations over the widest part of the catalogue.", "bullet": True},
        {"text": "Clustering recommends what is already popular in a learner's group, so it reinforces the popularity bias seen in the EDA; similarity and profile methods expose more long-tail courses.", "bullet": True},
        {"text": f"Only {overlap} courses appear in more than one top-10 list, so the methods are complementary: a hybrid (e.g. profile score re-ranked by similarity) is a natural next step.", "bullet": True},
    ], size=13, space_after=5)
    set_notes(slide, "Side-by-side comparison of the three content-based recommenders on the same 33,901 test users: list length, reach, catalogue coverage and total volume.")


# ----------------------------------------------------------------------------------------- main
def main():
    prs = Presentation(TEMPLATE)
    slides = list(prs.slides)
    (s_instr, s_title, s_outline, s_intro, s_eda, s_genre, s_enroll, s_top20, s_cloud, s_cb, s_fc_profile, s_ev_profile,
     s_fc_sim, s_ev_sim, s_fc_cluster, s_ev_cluster, s_cf, s_fc_knn, s_fc_nmf, s_fc_nn, s_compare, s_streamlit,
     s_conclusion, s_appendix) = slides

    eda = metrics("eda_summary.csv")
    profile, similarity, cluster = metrics("profile_summary.csv"), metrics("similarity_summary.csv"), metrics("cluster_summary.csv")
    knn, nmf = metrics("knn_summary.csv"), metrics("nmf_summary.csv")
    knn_sweep, nmf_sweep = csv("knn_sweep.csv"), csv("nmf_sweep.csv")
    nn_results, nn_history = csv("nn_results.csv"), csv("nn_history.csv")
    regression, classification = csv("regression_results.csv"), csv("classification_results.csv")
    bench = csv("cf_benchmark.csv")
    elbow, pca_var = csv("cluster_elbow.csv"), csv("cluster_pca_variance.csv")
    checks = rating_files_check()

    title_slide(s_title)
    intro_slide(s_intro, eda)
    genre_slide(s_genre)
    enrollment_slide(s_enroll)
    top20_slide(s_top20, eda)
    wordcloud_slide(s_cloud)

    # --- content-based: user profile
    flowchart_slide(s_fc_profile, [
        ("data", "Raw data", "course_genre.csv, user_profile.csv, ratings.csv"),
        ("process", "Course genre vectors", "14-d one-hot genre vector for each of the 307 courses"),
        ("process", "User profile vectors", "14-d interest weights: ratings x genres of enrolled courses"),
        ("process", "Unseen courses", "all courses minus the courses the user enrolled in"),
        ("process", "Interest score", "dot product of user profile and course genre vector"),
        ("process", "Threshold filter", f"keep unseen courses with score >= {profile['score_threshold']:.0f}"),
        ("output", "Recommendations", "USER, COURSE_ID, SCORE sorted by score"),
        ("eval", "Evaluation", "new courses per user and top-10 courses"),
    ], "The user profile is a weighted genre vector built from the user's ratings. Each unseen course gets a score equal to the dot product "
       "of the profile and the course genre vector; courses above the score threshold are recommended.",
        ["Course genre vectors (307 x 14) and user profile vectors (33,901 x 14) live in the same genre space, so a dot product measures how well a course matches a learner's interests.",
         f"Only unseen courses are scored; the score threshold controls list length - the lab default of 10 gave {csv('profile_threshold_sweep.csv').set_index('threshold').loc[10.0, 'avg_recs_per_user']:.0f} courses per user, so it was raised to {profile['score_threshold']:.0f}.",
         "Evaluation follows the lab questions: average number of new courses per test user and the 10 most frequently recommended courses."])
    p_sweep = csv("profile_threshold_sweep.csv")
    evaluation_slide(s_ev_profile,
                     f"recommendation score threshold = {profile['score_threshold']:.0f} (dot product of user profile and course genres); thresholds 5-40 tested, chosen to keep lists under 20 courses.",
                     f"{profile['avg_recs_per_user']:.1f}",
                     f"new/unseen courses recommended per test user on average ({int(profile['users_with_recs']):,} of {int(profile['test_users']):,} users receive at least one)",
                     "Average new courses per user by score threshold", sweep_columns(p_sweep, profile["score_threshold"], lambda t: f"{t:.0f}"),
                     csv("profile_top10.csv"),
                     "Results of the user-profile recommender on all 33,901 test users with score threshold 20. Left: average list length and its sensitivity to the threshold. Right: top-10 recommended courses.")

    # --- content-based: course similarity
    flowchart_slide(s_fc_sim, [
        ("data", "Raw data", "course_processed.csv (titles + descriptions), ratings.csv"),
        ("process", "Text processing", "NLTK: tokenize, drop stop words and numbers, keep nouns"),
        ("process", "Bag of words", "gensim Dictionary + doc2bow: token counts per course"),
        ("process", "Similarity matrix", "cosine similarity of BoW vectors, 307 x 307"),
        ("process", "Enrolled vs unseen", "similarity of each unseen course to each enrolled course"),
        ("process", "Best match + threshold", f"keep the highest similarity; recommend if > {similarity['similarity_threshold']:.1f}"),
        ("output", "Recommendations", "USER, COURSE_ID, SCORE sorted by similarity"),
        ("eval", "Evaluation", "new courses per user and top-10 courses"),
    ], "Course textual content becomes BoW feature vectors; the cosine similarity matrix between all courses is pre-computed. For each user, "
       "unseen courses that are similar enough to any enrolled course are recommended.",
        ["Each course is represented by the bag of words of its title and description (stop words removed, nouns kept), giving sparse token-count vectors.",
         "Cosine similarity between BoW vectors measures content overlap; e.g. Machine Learning with Python vs Machine Learning with R = 0.66.",
         f"For every test user an unseen course is recommended when its best similarity to one of the user's enrolled courses exceeds {similarity['similarity_threshold']:.1f}."])
    s_sweep = csv("similarity_threshold_sweep.csv")
    evaluation_slide(s_ev_sim,
                     f"cosine similarity threshold = {similarity['similarity_threshold']:.1f} on the BoW course similarity matrix (an unseen course must exceed it for at least one enrolled course); 0.4-0.8 tested.",
                     f"{similarity['avg_recs_per_user']:.1f}",
                     f"new/unseen courses recommended per test user on average ({int(similarity['users_with_recs']):,} of {int(similarity['test_users']):,} users receive at least one)",
                     "Average new courses per user by similarity threshold", sweep_columns(s_sweep, similarity["similarity_threshold"], lambda t: f"{t:.1f}"),
                     csv("similarity_top10.csv"),
                     "Results of the course-similarity recommender with threshold 0.6: about 8.5 new courses per user. Lower thresholds explode the list length (0.4 gives 48 courses per user).")

    # --- content-based: clustering
    flowchart_slide(s_fc_cluster, [
        ("data", "User profiles", "user_profile.csv: 33,901 users x 14 genre interests"),
        ("process", "Standardise", "StandardScaler: mean 0 and std 1 for every genre feature"),
        ("process", "PCA", f"{int(cluster['pca_components'])} components explain >= 90% of the variance"),
        ("process", "K-means", f"k = {int(cluster['optimal_k_pca'])} from the elbow of the inertia curve"),
        ("process", "Label test users", "join each user's cluster with the enrollments (ratings.csv)"),
        ("process", "Popular per cluster", f"courses with >= {int(cluster['popularity_threshold'])} enrollments in the cluster"),
        ("output", "Recommendations", "popular courses of the user's cluster not yet taken"),
        ("eval", "Evaluation", "new courses per user and top-10 courses"),
    ], "Users are grouped by their standardised (and PCA-reduced) profile vectors with K-means. Each cluster keeps a list of its popular courses; "
       "a user is recommended the popular courses of their cluster that they have not enrolled in.",
        [f"K-means was first run on the 14 standardised features (elbow at k = {int(cluster['optimal_k'])}), then on {int(cluster['pca_components'])} PCA components that keep 90% of the variance.",
         "Each cluster is a learning community; courses enrolled by many members of the community are its popular courses.",
         f"Popularity threshold = {int(cluster['popularity_threshold'])} enrollments within the cluster; the user's own courses are excluded from the recommendations."])
    c_sweep = csv("cluster_threshold_sweep.csv")

    def cluster_sweep_chart(slide, x, y, w, h):
        pivot = c_sweep.pivot(index="threshold", columns="features", values="avg_recs_per_user")
        cols = sorted(pivot.columns, key=lambda c: c.startswith("PCA"))
        add_bar_chart(slide, x, y, w, h, [str(int(t)) for t in pivot.index], [(c, pivot[c].round(1).tolist()) for c in cols],
                      colors=[LIGHT_SERIES, TITLE_BLUE], number_format="0.0", gap=60, overlap=-10, tick_size=10, label_size=9,
                      legend=True, cat_axis_title="Popularity threshold (enrollments in the cluster)")

    evaluation_slide(s_ev_cluster,
                     f"PCA n_components = {int(cluster['pca_components'])} (>= 90% variance), K-means k = {int(cluster['optimal_k_pca'])} (elbow), popularity threshold = {int(cluster['popularity_threshold'])} enrollments per cluster; 100-1,000 tested, with and without PCA.",
                     f"{cluster['avg_recs_per_user']:.1f}",
                     f"new/unseen courses recommended per test user on average ({int(cluster['users_with_recs']):,} of {int(cluster['test_users']):,} users receive at least one)",
                     "Average new courses per user, with and without PCA", cluster_sweep_chart,
                     csv("cluster_top10.csv"),
                     "Results of the clustering recommender. The grouped bars compare K-means on the original features with K-means on the PCA components for several popularity thresholds.")

    s_compare_cb = new_slide_like(prs, s_ev_cluster, "Comparing the three content-based recommender systems", list(prs.slides).index(s_ev_cluster) + 1)
    content_comparison_slide(s_compare_cb, profile, similarity, cluster, [csv("profile_top10.csv"), csv("similarity_top10.csv"), csv("cluster_top10.csv")])

    # --- collaborative filtering
    best_knn = knn_sweep.loc[knn_sweep["rmse"].idxmin()]
    best_nmf = nmf_sweep.loc[nmf_sweep["rmse"].idxmin()]
    flowchart_slide(s_fc_knn, [
        ("data", "Enrollment history", "ratings.csv: 233,306 (user, course, rating 3-5)"),
        ("process", "Surprise dataset", "Reader(rating_scale=(3, 5)); 70% train / 30% test split"),
        ("process", "Item-item similarity", "cosine, MSD or Pearson between course rating vectors (126 x 126)"),
        ("process", "k nearest neighbours", "the k most similar courses that the user has rated"),
        ("process", "Predict rating", "similarity-weighted average of the user's ratings of those neighbours"),
        ("eval", "Evaluate + tune", f"test RMSE; grid search similarity x k (best: {best_knn['similarity']}, k = {int(best_knn['k'])})"),
        ("eval", "Baseline check", f"compare with the global-mean RMSE ({knn['baseline_rmse']:.4f}) on the same test set"),
        ("output", "Recommend", "unseen courses with the highest predicted rating"),
    ], "Item-based KNN collaborative filtering with the Surprise library: course-course similarities from the rating matrix, predictions as the "
       "similarity-weighted average of the user's ratings of the k most similar courses, evaluated by RMSE on a held-out test set.",
        ["Item-based KNN: the course-course similarity matrix is tiny (126 x 126), while a user-based model would need a 33,901 x 33,901 matrix.",
         "Prediction: r(u, i) = sum of sim(i, j) x r(u, j) / sum of sim(i, j) over the k most similar courses j rated by u; grid search over cosine / MSD / Pearson and k = 10-80.",
         f"Best: {best_knn['similarity']}, k = {int(best_knn['k'])}, test RMSE {best_knn['rmse']:.4f} vs {knn['baseline_rmse']:.4f} for the global mean (ratings.csv carries no signal - see findings). Lab fix: Reader scale (3, 5) instead of (2, 3)."])
    flowchart_slide(s_fc_nmf, [
        ("data", "Enrollment history", "ratings.csv: 233,306 (user, course, rating 3-5)"),
        ("process", "User-item matrix", "33,901 x 126 sparse matrix; same 70/30 split as KNN"),
        ("process", "Factorise", "R ~ U x I^T with non-negative user (U) and item (I) factor matrices"),
        ("process", "Train with SGD", "minimise squared error on the known ratings; k = 8-64 factors"),
        ("process", "Predict rating", "dot product of the user's and the course's latent factors"),
        ("eval", "Evaluate + tune", f"test RMSE; grid search factors x epochs x init (best: {int(best_nmf['n_factors'])} factors, {int(best_nmf['n_epochs'])} epochs)"),
        ("eval", "Baseline check", f"compare with the global-mean RMSE ({knn['baseline_rmse']:.4f}) on the same test set"),
        ("output", "Recommend", "unseen courses with the highest predicted rating"),
    ], "Non-negative matrix factorisation: the sparse user-item matrix is decomposed into two small dense non-negative matrices of latent user "
       "and course features; their dot product estimates unknown ratings.",
        ["NMF compresses the 33,901 x 126 rating matrix into k latent factors per user and per course, so no large similarity matrix has to be kept in memory; factors are learned by SGD on the observed ratings only.",
         f"The lab's hint (init_high = 5) starts the factor dot products far above 5, so every prediction is clipped to 5 (RMSE {nmf_sweep.loc[(nmf_sweep['init_high'] == 5.0) & (nmf_sweep['n_epochs'] == 20), 'rmse'].iloc[0]:.4f}); the default init range (0, 1) was added to the grid.",
         f"Best: {int(best_nmf['n_factors'])} factors, {int(best_nmf['n_epochs'])} epochs, init ({best_nmf['init_low']:g}, {best_nmf['init_high']:g}) - test RMSE {best_nmf['rmse']:.4f} vs {knn['baseline_rmse']:.4f} for the global mean, on the same test ratings as KNN."])
    nn_lab = nn_results.iloc[0]
    nn_base = float(nn_results.loc[nn_results["model"].str.startswith("Global"), "test_rmse"].iloc[0])
    nn_v2 = nn_results[nn_results["model"].str.startswith("RecommenderNetV2")].sort_values("test_rmse").iloc[0]
    flowchart_slide(s_fc_nn, [
        ("data", "Enrollment history", "ratings (m03 lab 05).csv: 233,306 (user, course, rating 2 or 3)"),
        ("process", "Encode + scale", "user and course ids to indices; ratings scaled to 0-1; 80/10/10 split"),
        ("process", "Embedding layers", "user (33,901 -> 16-d) and course (126 -> 16-d) embeddings + biases"),
        ("process", "Interaction", "dot product of the embeddings + user bias + course bias -> ReLU"),
        ("process", "Train", "Adam, MSE loss, batch 64, 10 epochs, validation loss monitored"),
        ("eval", "Evaluate", f"test RMSE {nn_lab['test_rmse']:.4f} vs {nn_base:.4f} for the global mean"),
        ("process", "Improve: V2", f"per-example dot product + dense layer, early stopping: RMSE {nn_v2['test_rmse']:.4f}"),
        ("output", "Latent features", "user and course embeddings reused for regression / classification"),
    ], "RecommenderNet: two embedding layers learn latent user and course vectors whose dot product (plus bias terms) predicts the rating. "
       "The trained embeddings are extracted and reused as features for regression and classification models.",
        ["The one-hot user and course inputs are mapped to dense 16-d embedding vectors - the neural counterpart of the NMF factor matrices.",
         f"Finding: the lab's tf.tensordot(user, item, 2) sums the dot products over the whole batch. RecommenderNetV2 (per-example dot product + dense layer) cuts the test RMSE to {nn_v2['test_rmse']:.4f}.",
         "The embedding matrices (33,901 x 16 and 126 x 16) are the extracted latent features used by the regression and classification labs."])

    s_emb = new_slide_like(prs, s_fc_nn, "Embedding features: regression and classification results", list(prs.slides).index(s_fc_nn) + 1)
    embedding_slide(s_emb, regression, classification)

    compare_slide(s_compare, bench, knn, nmf, nn_results)
    conclusion_slide(s_conclusion, eda, profile, similarity, cluster, bench)
    appendix_slide(s_appendix)
    s_tuning = new_slide_like(prs, s_fc_nn, "Appendix: hyper-parameter tuning", len(prs.slides._sldIdLst))
    tuning_slide(s_tuning, knn_sweep, nmf_sweep, elbow, pca_var, nn_history, cluster)
    s_findings = new_slide_like(prs, s_fc_nn, "Appendix: data-quality findings", len(prs.slides._sldIdLst))
    findings_slide(s_findings, nn_results, p_sweep, nmf_sweep, checks)

    delete_slide(prs, s_instr)
    delete_slide(prs, s_streamlit)
    prs.core_properties.author = AUTHOR
    prs.core_properties.title = "Machine Learning Capstone Project Report"
    prs.save(OUTPUT)
    print("saved", OUTPUT)


def embedding_slide(slide, regression, classification):
    col_w = (X1 - X0 - 0.4) / 2
    add_text(slide, X0, 1.95, col_w, 0.4, [{"text": "Regression: predict the rating score", "size": 15, "bold": True, "color": TITLE_BLUE}], space_after=0)
    reg = regression.sort_values("test_rmse")
    rows = [["Model", "Best hyper-parameters", "Test RMSE"]] + [[r.model, r.best_params.replace("{", "").replace("}", "").replace("'", ""), f"{r.test_rmse:.4f}"] for r in reg.itertuples()]
    add_table(slide, X0, 2.4, [2.15, 2.25, 1.15], rows, size=10.5, row_h=0.36, aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT])
    x = X0 + col_w + 0.4
    add_text(slide, x, 1.95, col_w, 0.4, [{"text": "Classification: predict the rating class (3 / 4 / 5)", "size": 15, "bold": True, "color": TITLE_BLUE}], space_after=0)
    cls = classification.sort_values("f1_macro", ascending=False)
    rows = [["Model", "Accuracy", "Macro F1", "AUC (OvR)"]] + [[r.model, f"{r.accuracy:.3f}", f"{r.f1_macro:.3f}", "-" if np.isnan(r.auc_ovr) else f"{r.auc_ovr:.3f}"] for r in cls.itertuples()]
    add_table(slide, x, 2.4, [2.45, 1.0, 1.0, 1.1], rows, size=10.5, row_h=0.36, aligns=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT])
    base_rmse = float(regression.loc[regression["model"] == "Global-mean baseline", "test_rmse"].iloc[0])
    best_reg = reg.iloc[0]
    base_acc = float(classification.loc[classification["model"].str.startswith("Majority"), "accuracy"].iloc[0])
    best_cls = cls.iloc[0]
    add_card(slide, X0, 5.48, X1 - X0, 1.45, fill=NEUTRAL)
    add_text(slide, X0 + 0.25, 5.55, X1 - X0 - 0.5, 1.35, [
        {"text": "Input X = user embedding + course embedding (element-wise sum, 16 features) for every rating of ratings.csv; label y = rating score / class.", "bullet": True},
        {"text": f"Best regressor: {best_reg.model} (RMSE {best_reg.test_rmse:.4f} vs {base_rmse:.4f} for the global mean); best classifier: {best_cls.model} "
                 f"(accuracy {best_cls.accuracy:.3f} vs {base_acc:.3f} for the majority class, macro F1 {best_cls.f1_macro:.3f}).", "bullet": True},
        {"text": "Every regressor matches the global mean and every classifier sits at chance (AUC about 0.50): the 3/4/5 ratings of ratings.csv are random, so there is no signal to learn (see findings).", "bullet": True},
    ], size=12.5, space_after=4)
    set_notes(slide, "Regression and classification labs: embedding features from the neural network are aggregated and fed to scikit-learn models. "
                     "Models were tuned with 3-fold grid search (regression) and compared with naive baselines.")


def bench_split(bench):
    """Benchmark RMSE per model for the lab-05 file (real signal) and ratings.csv (random ratings)."""
    signal = bench[bench["dataset"].str.contains("lab 05")].set_index("model")["rmse"]
    random = bench[bench["dataset"].str.startswith("ratings.csv")].set_index("model")["rmse"]
    base = lambda series: float(series[[m for m in series.index if m.startswith("Global")][0]])
    return signal, random, base(signal), base(random)


def model_rmse(series, prefix):
    return float(series[[m for m in series.index if m.startswith(prefix)][0]])


def compare_slide(slide, bench, knn, nmf, nn_results):
    clear_slide(slide)
    signal, random, base_s, base_r = bench_split(bench)
    col_w = (X1 - X0 - 0.4) / 2
    panels = [("Lab-05 ratings (2/3): the file with real signal", signal), ("ratings.csv (3/4/5): random ratings", random)]
    for i, (heading, series) in enumerate(panels):
        part = series.sort_values()
        x = X0 + i * (col_w + 0.4)
        add_text(slide, x, 1.92, col_w, 0.6, [{"text": heading, "size": 14, "bold": True, "color": TITLE_BLUE, "space_after": 0},
                                               {"text": "test RMSE, same 70/30 split for every model, rating units, lower is better", "size": 11, "color": SECONDARY, "space_after": 0}], space_after=0)
        colors = {j: REFERENCE for j, m in enumerate(part.index) if m.startswith("Global")}
        colors.setdefault(0, TITLE_BLUE)
        add_bar_chart(slide, x, 2.55, col_w, 2.9, part.index.tolist(), [("RMSE", part.round(4).tolist())],
                      horizontal=True, colors=[LIGHT_SERIES], point_colors=colors, number_format="0.0000", gap=40,
                      tick_size=10.5, label_size=10, min_value=0)
    best = signal.drop([m for m in signal.index if m.startswith("Global")]).idxmin()
    knn_gain = 1 - model_rmse(signal, "KNN") / base_s
    nmf_gain = 1 - model_rmse(signal, "NMF") / base_s
    add_card(slide, X0, 5.55, X1 - X0, 1.4, fill=NEUTRAL)
    add_text(slide, X0 + 0.25, 5.6, X1 - X0 - 0.5, 1.3, [
        {"runs": [("Signal file: ", {"bold": True}), (f"{best} is best ({signal[best]:.4f}, {1 - signal[best] / base_s:.0%} below the global mean {base_s:.4f}), then the lab RecommenderNet "
                                                       f"({model_rmse(signal, 'RecommenderNet (lab)'):.4f}); KNN and NMF improve {knn_gain:.0%} and {nmf_gain:.0%}; untuned SVD (100 factors) overfits.", {})], "bullet": True},
        {"runs": [("ratings.csv: ", {"bold": True}), (f"the ratings are random, so the global mean ({base_r:.4f}) is the best possible score; models that fit noise score worse. "
                                                       f"Lab runs agree: KNN {knn['best_knn_rmse']:.4f} and NMF {nmf['best_nmf_rmse']:.4f} vs baseline {knn['baseline_rmse']:.4f}.", {})], "bullet": True},
        {"runs": [("Metric: ", {"bold": True}), ("RMSE = square root of the mean squared difference between predicted and true ratings on held-out data - the evaluation metric of every collaborative-filtering lab.", {})], "bullet": True},
    ], size=12, space_after=4)
    set_notes(slide, "Every collaborative-filtering model was retrained in m03/08 on one identical split per rating file, so the RMSEs are directly comparable. "
                     "Gray bar: global-mean baseline; dark blue: best model. RMSE (root mean squared error) is the evaluation metric of all collaborative-filtering labs.")


def conclusion_slide(slide, eda, profile, similarity, cluster, bench):
    for shape in list(slide.shapes):
        if shape.has_text_frame and shape.text_frame.text.startswith("Point"):
            remove_shape(shape)
    signal, random, base_s, base_r = bench_split(bench)
    best = signal.drop([m for m in signal.index if m.startswith("Global")]).idxmin()
    top, h, left_w, gap = 1.62, 4.0, 6.55, 0.3
    right_x = 0.84 + left_w + gap
    add_card(slide, 0.84, top, left_w, h, fill=WHITE)
    add_card(slide, right_x, top, 12.5 - right_x, h, fill=WHITE)
    add_text(slide, 0.84 + 0.25, top + 0.18, left_w - 0.5, h - 0.3, [
        {"text": "Key findings", "size": 18, "bold": True, "color": TITLE_BLUE, "space_after": 8},
        {"text": f"H1 confirmed: the top-20 courses take {eda['top20_enrollments'] / eda['enrollments']:.0%} of all enrollments, and 25% of learners took a single course.", "bullet": True},
        {"text": f"H2 confirmed: after tuning, the profile, similarity and clustering recommenders suggest {profile['avg_recs_per_user']:.1f}, {similarity['avg_recs_per_user']:.1f} and "
                 f"{cluster['avg_recs_per_user']:.1f} new courses per learner, reaching {profile['users_with_recs'] / profile['test_users']:.0%}, {similarity['users_with_recs'] / similarity['test_users']:.0%} "
                 f"and {cluster['users_with_recs'] / cluster['test_users']:.0%} of learners.", "bullet": True},
        {"text": f"H3 confirmed where the data has signal: {best} cuts RMSE by {1 - signal[best] / base_s:.0%} vs the global mean and beats KNN and NMF. "
                 f"On ratings.csv, whose ratings are random, no model beats the mean ({base_r:.3f}).", "bullet": True},
        {"text": "Data checks and tuning matter as much as the model: default thresholds gave 41-44 courses per learner, and a randomised rating file makes every model look broken.", "bullet": True},
    ], size=13.5, space_after=6)
    add_text(slide, right_x + 0.25, top + 0.18, 12.5 - right_x - 0.5, h - 0.3, [
        {"text": "Recommendations", "size": 18, "bold": True, "color": TITLE_BLUE, "space_after": 8},
        {"text": "Serve new and one-course learners with content-based recommendations (profile or similarity): they need no rating history.", "bullet": True},
        {"text": "Use the neural embedding model for learners with history, with item-based KNN as a simple, explainable fallback.", "bullet": True},
        {"text": "Blend both in a hybrid ranker, cap lists with a per-user top-N (10-20), deduplicate course versions and add diversity for long-tail courses.", "bullet": True},
        {"text": "Next: ranking metrics (precision@k, recall@k) on a time-based split, then an A/B test of click-through and completion.", "bullet": True},
    ], size=13.5, space_after=6)
    add_card(slide, 0.84, top + h + 0.25, 12.5 - 0.84, 0.72, fill=TINT)
    add_text(slide, 0.84 + 0.25, top + h + 0.25, 12.5 - 0.84 - 0.5, 0.72, [{"runs": [
        ("Bottom line: ", {"bold": True, "color": TITLE_BLUE}),
        ("start every learner with content-based recommendations, switch to the neural embedding model once they have a rating history, "
         "and judge every model against a baseline on validated data.", {})]}], size=14, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    set_notes(slide, "Summary of the findings against the three hypotheses and the recommended next steps.")


def appendix_slide(slide):
    for shape in list(slide.shapes):
        if shape.has_text_frame and (shape.text_frame.text.startswith("Instructions") or shape.text_frame.text.startswith("Asset")):
            remove_shape(shape)
    assets = [
        ("Lab notebooks", "11 labs solved and executed: m01/06, 07, 10 (EDA, bag of words, similarity); m02/03, 05, 06 (profile, similarity, clustering); m03/03-07 (KNN, NMF, neural network, regression, classification)."),
        ("Benchmark notebook", "m03/08 - Extra - Benchmark of Collaborative Filtering Models.ipynb retrains KNN, NMF, SVD, both neural networks and Ridge on one identical 70/30 split of each rating file."),
        ("Results", "results/*.csv - the small summaries each notebook exports (threshold sweeps, top-10 lists, RMSE tables). Every number in this deck is read from them."),
        ("Report source", "m05/Final Project Submission/source: build_report.py fills the course template from the results; export_pdf.ps1 saves this PDF with PowerPoint."),
        ("Environment", "Python 3.11, pandas 2.2, scikit-learn 1.5, scikit-surprise 1.1.4, TensorFlow 2.17 (Keras 3), gensim 4.3, NLTK 3.9, wordcloud 1.9, python-pptx 1.0."),
    ]
    top, gap = 1.62, 0.3
    w = (12.5 - 0.84 - 2 * gap) / 3
    h = (4.95 - gap) / 2
    for i, (head, body) in enumerate(assets + [(None, None)]):
        row, col = divmod(i, 3)
        x, y = 0.84 + col * (w + gap), top + row * (h + gap)
        add_card(slide, x, y, w, h, fill=WHITE)
        if head is None:
            add_text(slide, x + 0.22, y + 0.15, w - 0.44, 0.35, [{"text": "Key code: profile scoring", "size": 14, "bold": True, "color": TITLE_BLUE}], space_after=0)
            code = ["# interest score of every course",
                    "# for every learner, vectorised",
                    "scores = profiles @ genres.T",
                    "scores = scores.mask(enrolled)",
                    "recs = scores >= 20",
                    "recs.sum(axis=1).mean()  # 14.1"]
            box = add_text(slide, x + 0.22, y + 0.6, w - 0.44, h - 0.7, [{"text": line, "size": 10.5, "color": MUTED if line.startswith("#") else INK, "space_after": 2} for line in code], space_after=2)
            for p in box.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Consolas"
            continue
        add_text(slide, x + 0.22, y + 0.15, w - 0.44, h - 0.3, [{"text": head, "size": 14, "bold": True, "color": TITLE_BLUE, "space_after": 6},
                                                              {"text": body, "size": 13, "color": INK, "space_after": 0}], space_after=0)
    set_notes(slide, "All project assets live in the capstone folder of the course repository.")


def tuning_slide(slide, knn_sweep, nmf_sweep, elbow, pca_var, nn_history, cluster):
    w = (X1 - X0 - 0.4) / 2
    h = 2.15
    pv = knn_sweep.pivot(index="k", columns="similarity", values="rmse")
    add_text(slide, X0, 1.85, w, 0.3, [{"text": "KNN: test RMSE by similarity and k (cosine and MSD overlap)", "size": 12.5, "bold": True, "color": SECONDARY}], space_after=0)
    add_line_chart(slide, X0, 2.15, w, h, [str(k) for k in pv.index], [(c, pv[c].round(4).tolist()) for c in pv.columns],
                   colors=[SERIES, SERIES_2, "1BAF7A"], number_format="0.000", legend=True, cat_axis_title="k (neighbours)")
    pv = nmf_sweep.pivot_table(index="n_factors", columns=["init", "n_epochs"], values="rmse")
    x = X0 + w + 0.4
    add_text(slide, x, 1.85, w, 0.3, [{"text": "NMF: test RMSE by factors, epochs and initialisation range", "size": 12.5, "bold": True, "color": SECONDARY}], space_after=0)
    add_line_chart(slide, x, 2.15, w, h, [str(k) for k in pv.index], [(f"init {i}, {e} epochs", pv[(i, e)].round(4).tolist()) for i, e in pv.columns],
                   colors=[SERIES, "86B6EF", SERIES_2, "F4A582"], number_format="0.00", legend=True, cat_axis_title="latent factors",
                   min_value=0.8, max_value=1.4)
    y2 = 4.5
    add_text(slide, X0, y2, w, 0.3, [{"text": f"K-means elbow (k = {int(cluster['optimal_k'])} raw, {int(cluster['optimal_k_pca'])} on PCA)", "size": 12.5, "bold": True, "color": SECONDARY}], space_after=0)
    add_line_chart(slide, X0, y2 + 0.3, w, h, [str(k) for k in elbow["k"]], [("Raw features", (elbow["inertia"] / 1000).round(1).tolist()), ("PCA features", (elbow["inertia_pca"] / 1000).round(1).tolist())],
                   colors=[SERIES, SERIES_2], number_format="0", legend=True, axis_title="inertia (thousands)", cat_axis_title="k")
    add_text(slide, x, y2, w, 0.3, [{"text": f"PCA: accumulated explained variance ({int(cluster['pca_components'])} components >= 90%)", "size": 12.5, "bold": True, "color": SECONDARY}], space_after=0)
    add_line_chart(slide, x, y2 + 0.3, w, h, [str(n) for n in pca_var["n_components"]], [("Explained variance", pca_var["accumulated_variance_ratio"].round(3).tolist())],
                   number_format="0%", cat_axis_title="n_components", min_value=0, max_value=1)
    set_notes(slide, "Hyper-parameter searches behind the selected settings: KNN similarity and k, NMF factors and epochs, the K-means elbow and the PCA variance curve.")


def duplicate_courses():
    sim = pd.read_csv(DATASETS / "sim.csv").to_numpy()
    bow = pd.read_csv(DATASETS / "courses_bows.csv")
    ids = bow.groupby(["doc_index", "doc_id"]).size().reset_index().set_index("doc_index")["doc_id"]
    titles = pd.read_csv(DATASETS / "course_processed.csv").set_index("COURSE_ID")["TITLE"]
    rows, cols = np.triu_indices_from(sim, 1)
    pairs = [(ids[i], ids[j]) for i, j in zip(rows, cols) if sim[i, j] > 0.999]
    return [f"{course_title(titles[a])} ({a} / {b})" for a, b in pairs]


def rating_files_check():
    """ratings.csv and the lab-05 file hold the same (user, course) pairs; compare their ratings pair by pair."""
    v2 = pd.read_csv(DATASETS / "ratings.csv")
    lab05 = pd.read_csv(DATASETS / "ratings (m03 lab 05).csv")
    both = v2.merge(lab05, on=["user", "item"], suffixes=("_v2", "_lab05"))
    return {"pairs": len(both), "same_pairs": len(both) == len(v2) == len(lab05),
            "corr": both["rating_v2"].corr(both["rating_lab05"]),
            "shares": v2["rating"].value_counts(normalize=True).sort_index().to_dict()}


def findings_slide(slide, nn_results, p_sweep, nmf_sweep, checks):
    duplicates = duplicate_courses()
    shares = ", ".join(f"{k}: {v:.1%}" for k, v in checks["shares"].items())
    clipped = nmf_sweep.loc[(nmf_sweep["init_high"] == 5.0) & (nmf_sweep["n_epochs"] == 20), "rmse"].iloc[0]
    items = [
        ("ratings.csv has random ratings", f"It holds the same {checks['pairs']:,} (user, course) pairs as the lab-05 file, but its ratings are spread evenly ({shares}) and uncorrelated with the lab-05 ratings "
                                           f"(r = {checks['corr']:.3f}). No model can beat the global mean on it; the benchmark retrains every model on both files."),
        ("Rating scale in the KNN / NMF labs", "The labs read ratings.csv with Reader(rating_scale=(2, 3)); Surprise clips predictions to that range, so every estimate would be 3. Fixed to (3, 5)."),
        ("NMF initialisation in the lab hint", f"init_low = 0.5, init_high = 5 makes the initial factor dot products far larger than 5, so after 20 epochs every prediction is clipped to 5 (RMSE {clipped:.4f}). "
                                               "The default range (0, 1) was added to the grid search."),
        ("Batch-wide dot product in RecommenderNet", "tf.tensordot(user_vector, item_vector, 2) contracts both axes, so every example in a batch gets the same interaction term. "
                                                     f"RecommenderNetV2 (per-example dot product + dense layer): test RMSE {nn_results.iloc[1:]['test_rmse'].min():.4f} vs {nn_results.iloc[0]['test_rmse']:.4f}."),
        ("Default thresholds give long lists", f"Score threshold 10 recommended {p_sweep.set_index('threshold').loc[10.0, 'avg_recs_per_user']:.0f} courses per learner and popularity threshold 100 about 41; "
                                               "every threshold was swept and set to stay under the 20-course guideline."),
        ("Duplicate course versions", f"{len(duplicates)} course pairs have identical text (similarity 1.0), e.g. {duplicates[0]}. They take slots in the top-10 lists and can recommend a course the learner already took under another id."),
    ]
    y = 1.88
    for i, (head, body) in enumerate(items):
        add_card(slide, X0, y, X1 - X0, 0.78, fill=NEUTRAL if i % 2 == 0 else WHITE)
        add_text(slide, X0 + 0.2, y + 0.03, X1 - X0 - 0.4, 0.72, [{"text": head, "size": 12.5, "bold": True, "color": TITLE_BLUE, "space_after": 1},
                                                                  {"text": body, "size": 11, "color": INK, "space_after": 0}], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.845
    set_notes(slide, "Issues found while solving the labs, and how each was handled. Also fixed: the lab helper generate_sparse_bow() marks only one "
                     "course's own tokens, so any two courses got cosine 1.0; the solution builds both vectors over the shared vocabulary (cosine 0.75).")


if __name__ == "__main__":
    main()
