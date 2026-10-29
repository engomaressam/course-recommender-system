"""Small python-pptx helpers for filling the capstone template: text, cards, charts, tables, flowcharts."""
import copy
import math

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Abadi"
TITLE_BLUE = "0B49CB"   # the template's title colour
SERIES = "2A78D6"       # single-series marks
SERIES_2 = "EB6834"     # second categorical slot
DARK_BLUE = "104281"
INK = "262626"
SECONDARY = "52514E"
MUTED = "898781"
GRID = "E1E0D9"
AXIS = "C3C2B7"
REFERENCE = "B4B2AA"    # baseline / reference marks
TINT = "EEF4FD"         # light blue card fill
NEUTRAL = "F4F4F1"      # neutral card fill
WHITE = "FFFFFF"

X0, X1 = 0.92, 12.42    # content left / right edge (inches), aligned with the title placeholder
Y0, Y1 = 1.85, 6.95     # content top / bottom


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


# ----------------------------------------------------------------------------------------- slides
def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
            return


def move_slide(prs, slide, new_index):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            sld_ids.remove(sld_id)
            sld_ids.insert(new_index, sld_id)
            return


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def clear_slide(slide, keep_title=True):
    """Remove every shape except the title placeholder (instruction boxes, sample images, sample flowcharts)."""
    for shape in list(slide.shapes):
        is_title = shape.is_placeholder and shape.placeholder_format.type is not None and "TITLE" in str(shape.placeholder_format.type)
        if keep_title and is_title:
            continue
        remove_shape(shape)


def set_runs_text(shape, text):
    """Replace the text of a shape but keep the formatting of its first run."""
    p = shape.text_frame.paragraphs[0]
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def new_slide_like(prs, ref_slide, title, index):
    """Add a slide on ref_slide's layout with a copy of its title shape (keeps the template's title style)."""
    slide = prs.slides.add_slide(ref_slide.slide_layout)
    for shape in list(slide.shapes):
        remove_shape(shape)
    ref_title = ref_slide.shapes.title
    slide.shapes._spTree.append(copy.deepcopy(ref_title._element))
    set_runs_text(slide.shapes.title, title)
    move_slide(prs, slide, index)
    return slide


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ----------------------------------------------------------------------------------------- text
def _apply_run(run, size, bold=False, color=INK, italic=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = rgb(color)


def _bullet(paragraph, level=0, char="•", color=SERIES, indent=0.22):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(indent * (level + 1)))))
    pPr.set("indent", str(-Emu(Inches(indent))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def fill_text(tf, paragraphs, size=14, color=INK, align=PP_ALIGN.LEFT, space_after=6, line_spacing=None):
    """paragraphs: list of str | dict(text|runs, size, bold, color, bullet, level, align, space_after, italic).
    runs: list of (text, {bold, color, size, italic})."""
    tf.word_wrap = True
    for i, spec in enumerate(paragraphs):
        if isinstance(spec, str):
            spec = {"text": spec}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.space_after = Pt(spec.get("space_after", space_after))
        if line_spacing or spec.get("line_spacing"):
            p.line_spacing = spec.get("line_spacing", line_spacing)
        runs = spec.get("runs") or [(spec.get("text", ""), {})]
        for text, opts in runs:
            r = p.add_run()
            r.text = text
            _apply_run(r, opts.get("size", spec.get("size", size)), opts.get("bold", spec.get("bold", False)),
                       opts.get("color", spec.get("color", color)), opts.get("italic", spec.get("italic", False)))
        if spec.get("bullet"):
            _bullet(p, spec.get("level", 0))
    return tf


def add_text(slide, x, y, w, h, paragraphs, size=14, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             margin=0.0, space_after=6, name=None, line_spacing=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    fill_text(tf, paragraphs if isinstance(paragraphs, list) else [paragraphs], size, color, align, space_after, line_spacing)
    if name:
        box.name = name
    return box


def add_card(slide, x, y, w, h, fill=TINT, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    card = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        card.adjustments[0] = radius
    card.fill.solid()
    card.fill.fore_color.rgb = rgb(fill)
    if line:
        card.line.color.rgb = rgb(line)
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    card.shadow.inherit = False
    card.text_frame.text = ""
    if name:
        card.name = name
    return card


def add_stat(slide, x, y, w, h, value, label, fill=TINT, value_size=34, label_size=12, value_color=TITLE_BLUE):
    """A stat tile: big number with a short label below, on a tinted card."""
    add_card(slide, x, y, w, h, fill=fill)
    add_text(slide, x + 0.18, y + 0.12, w - 0.36, h - 0.24,
             [{"text": value, "size": value_size, "bold": True, "color": value_color, "space_after": 2},
              {"text": label, "size": label_size, "color": SECONDARY, "space_after": 0}],
             anchor=MSO_ANCHOR.MIDDLE)


def add_badge(slide, x, y, number, d=0.34, fill=TITLE_BLUE):
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    badge.fill.solid()
    badge.fill.fore_color.rgb = rgb(fill)
    badge.line.color.rgb = rgb(WHITE)
    badge.line.width = Pt(1.5)
    badge.shadow.inherit = False
    tf = badge.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_text(tf, [{"text": str(number), "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "space_after": 0}])
    return badge


# ----------------------------------------------------------------------------------------- connectors
def add_arrow(slide, x1, y1, x2, y2, color=SECONDARY, width=1.5):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return conn


# ----------------------------------------------------------------------------------------- flowchart
KIND_STYLE = {
    # kind: (shape, fill, line, title colour, detail colour)
    "data": (MSO_SHAPE.FLOWCHART_MAGNETIC_DISK, NEUTRAL, AXIS, INK, SECONDARY),
    "process": (MSO_SHAPE.ROUNDED_RECTANGLE, TINT, SERIES, INK, SECONDARY),
    "eval": (MSO_SHAPE.ROUNDED_RECTANGLE, WHITE, SERIES, INK, SECONDARY),
    "output": (MSO_SHAPE.FLOWCHART_DOCUMENT, TITLE_BLUE, TITLE_BLUE, WHITE, "DCE6FA"),
}


def add_flowchart(slide, steps, top=2.0, box_w=2.55, box_h=1.2, gap=0.42, row_gap=0.52, title_size=12.5, detail_size=10.5):
    """Snake layout: row 1 left->right, a down arrow on the right, row 2 right->left.
    steps: list of (kind, title, detail). Returns the bottom y of the chart."""
    per_row = math.ceil(len(steps) / 2)
    total_w = per_row * box_w + (per_row - 1) * gap
    left = X0 + (X1 - X0 - total_w) / 2
    boxes = []
    for i, (kind, title, detail) in enumerate(steps):
        row, col = divmod(i, per_row)
        if row == 1:
            col = per_row - 1 - col
        x, y = left + col * (box_w + gap), top + row * (box_h + row_gap)
        shape_type, fill, line, title_color, detail_color = KIND_STYLE[kind]
        shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(box_w), Inches(box_h))
        if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = 0.12
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1.25)
        if kind == "eval":
            shp.line.dash_style = 4  # dash
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.1)
        tf.margin_top = Inches(0.22 if shape_type == MSO_SHAPE.FLOWCHART_MAGNETIC_DISK else 0.06)
        tf.margin_bottom = Inches(0.12 if kind == "output" else 0.06)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_text(tf, [{"text": title, "size": title_size, "bold": True, "color": title_color, "align": PP_ALIGN.CENTER, "space_after": 2},
                       {"text": detail, "size": detail_size, "color": detail_color, "align": PP_ALIGN.CENTER, "space_after": 0}])
        add_badge(slide, x - 0.12, y - 0.14, i + 1)
        boxes.append((x, y, row))
    for (xa, ya, ra), (xb, yb, rb) in zip(boxes, boxes[1:]):
        if ra == rb == 0:
            add_arrow(slide, xa + box_w, ya + box_h / 2, xb, yb + box_h / 2)
        elif ra == rb == 1:
            add_arrow(slide, xa, ya + box_h / 2, xb + box_w, yb + box_h / 2)
        else:
            add_arrow(slide, xa + box_w / 2, ya + box_h, xb + box_w / 2, yb)
    return top + 2 * box_h + row_gap


# ----------------------------------------------------------------------------------------- charts
def _style_axis_text(axis, size, color):
    axis.tick_labels.font.size = Pt(size)
    axis.tick_labels.font.name = FONT
    axis.tick_labels.font.color.rgb = rgb(color)


def _cat_axis_skip(chart, skip):
    cat_ax = chart._chartSpace.find(".//" + qn("c:catAx"))
    el = etree.Element(qn("c:tickLblSkip"))
    el.set("val", str(skip))
    anchor = cat_ax.find(qn("c:noMultiLvlLbl"))
    if anchor is not None:
        anchor.addprevious(el)
    else:
        cat_ax.append(el)


def add_bar_chart(slide, x, y, w, h, categories, series, horizontal=False, colors=None, point_colors=None,
                  number_format="0", label_size=10, tick_size=10, gap=60, overlap=0, show_values=True,
                  show_value_axis=False, max_value=None, min_value=None, legend=False, axis_title=None,
                  cat_axis_title=None, tick_skip=None, label_position=XL_LABEL_POSITION.OUTSIDE_END):
    """series: list of (name, values). Single-series charts get no legend (the slide title names them)."""
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart_type = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gframe = slide.shapes.add_chart(chart_type, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = gframe.chart
    chart.has_title = False  # the slide or a heading above the chart names it
    if show_values and not show_value_axis and max_value is None:
        # headroom so the label of the longest bar is not clipped
        max_value = max(max(v for v in values if v is not None) for _, values in series) * (1.16 if horizontal else 1.12)
    chart.font.name = FONT
    chart.font.size = Pt(tick_size)
    chart.font.color.rgb = rgb(SECONDARY)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(tick_size)
        chart.legend.font.name = FONT
    plot = chart.plots[0]
    plot.gap_width = gap
    plot.overlap = overlap
    plot.vary_by_categories = False
    colors = colors or [SERIES, SERIES_2, "1BAF7A", "EDA100"]
    for i, ser in enumerate(plot.series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
        ser.format.line.fill.background()
        ser.invert_if_negative = False
    if point_colors:
        for idx, color in point_colors.items():
            pt = plot.series[0].points[idx]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(color)
    if show_values:
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = number_format
        labels.number_format_is_linked = False
        labels.position = label_position
        labels.font.size = Pt(label_size)
        labels.font.name = FONT
        labels.font.color.rgb = rgb(INK)
    value_axis, cat_axis = chart.value_axis, chart.category_axis
    if max_value is not None:
        value_axis.maximum_scale = max_value
    if min_value is not None:
        value_axis.minimum_scale = min_value
    value_axis.major_tick_mark = XL_TICK_MARK.NONE
    if show_value_axis:
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = rgb(GRID)
        value_axis.major_gridlines.format.line.width = Pt(0.75)
        value_axis.format.line.fill.background()
        value_axis.tick_labels.number_format = number_format
        value_axis.tick_labels.number_format_is_linked = False
        _style_axis_text(value_axis, tick_size, MUTED)
    else:
        value_axis.has_major_gridlines = False
        value_axis.visible = False
    cat_axis.major_tick_mark = XL_TICK_MARK.NONE
    cat_axis.format.line.color.rgb = rgb(AXIS)
    cat_axis.has_major_gridlines = False
    _style_axis_text(cat_axis, tick_size, SECONDARY)
    if horizontal:
        cat_axis.reverse_order = True
    if tick_skip:
        _cat_axis_skip(chart, tick_skip)
    for axis, title in ((value_axis, axis_title), (cat_axis, cat_axis_title)):
        if title:
            axis.has_title = True
            axis.axis_title.text_frame.text = title
            run = axis.axis_title.text_frame.paragraphs[0].runs[0]
            _apply_run(run, tick_size, color=MUTED)
    return chart


def add_line_chart(slide, x, y, w, h, categories, series, colors=None, number_format="0", tick_size=10,
                   axis_title=None, cat_axis_title=None, legend=False, min_value=None, max_value=None):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = gframe.chart
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(tick_size)
    chart.font.color.rgb = rgb(SECONDARY)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(tick_size)
    colors = colors or [SERIES, SERIES_2]
    for i, ser in enumerate(chart.plots[0].series):
        color = rgb(colors[i % len(colors)])
        ser.smooth = False
        ser.format.line.color.rgb = color
        ser.format.line.width = Pt(2)
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 7
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = color
        ser.marker.format.line.color.rgb = rgb(WHITE)
    value_axis, cat_axis = chart.value_axis, chart.category_axis
    if min_value is not None:
        value_axis.minimum_scale = min_value
    if max_value is not None:
        value_axis.maximum_scale = max_value
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = rgb(GRID)
    value_axis.major_gridlines.format.line.width = Pt(0.75)
    value_axis.format.line.fill.background()
    value_axis.major_tick_mark = XL_TICK_MARK.NONE
    value_axis.tick_labels.number_format = number_format
    value_axis.tick_labels.number_format_is_linked = False
    _style_axis_text(value_axis, tick_size, MUTED)
    cat_axis.major_tick_mark = XL_TICK_MARK.NONE
    cat_axis.format.line.color.rgb = rgb(AXIS)
    _style_axis_text(cat_axis, tick_size, SECONDARY)
    for axis, title in ((value_axis, axis_title), (cat_axis, cat_axis_title)):
        if title:
            axis.has_title = True
            axis.axis_title.text_frame.text = title
            _apply_run(axis.axis_title.text_frame.paragraphs[0].runs[0], tick_size, color=MUTED)
    return chart


# ----------------------------------------------------------------------------------------- tables
def _cell_border_bottom(cell, color=GRID, width_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for tag, visible in (("a:lnL", False), ("a:lnR", False), ("a:lnT", False), ("a:lnB", True)):
        ln = etree.Element(qn(tag))
        if visible:
            ln.set("w", str(int(width_pt * 12700)))
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        # borders must precede the cell fill
        first_fill = next((c for c in tcPr if c.tag in (qn("a:solidFill"), qn("a:noFill"))), None)
        if first_fill is not None:
            first_fill.addprevious(ln)
        else:
            tcPr.append(ln)


def add_table(slide, x, y, col_widths, rows, header=True, size=11, row_h=0.3, header_fill=TITLE_BLUE,
              aligns=None, zebra=NEUTRAL, bold_first_col=False):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(sum(col_widths)), Inches(row_h * n_rows))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1" if header else "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for j, wdt in enumerate(col_widths):
        tbl.columns[j].width = Inches(wdt)
    aligns = aligns or [PP_ALIGN.LEFT] * n_cols
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        is_header = header and i == 0
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if is_header else (zebra if (i % 2 == 0 and zebra) else WHITE))
            tf = cell.text_frame
            tf.word_wrap = True
            fill_text(tf, [{"text": str(value), "size": size, "bold": is_header or (bold_first_col and j == 0),
                            "color": WHITE if is_header else INK, "align": aligns[j], "space_after": 0}])
            if not is_header:
                _cell_border_bottom(cell)
    return tbl
